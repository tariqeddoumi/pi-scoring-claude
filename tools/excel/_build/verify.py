import json, sys, time, warnings, re
warnings.filterwarnings("ignore")
import formulas
X = "PI_Promotion_Modele_v4.xlsx"
vec = json.load(open("_build/vectors.json", encoding="utf-8"))
t0 = time.time()
xl = formulas.ExcelModel().loads(X).finish()
print(f"chargement: {time.time()-t0:.1f}s", flush=True)
fn = X.split("/")[-1]
def ref(sheet, cell): return f"'[{fn}]{sheet.upper()}'!{cell}"
# clés -> ligne dans Saisie (13..)
import openpyxl
wb = openpyxl.load_workbook(X)
ws = wb["Saisie"]
rowof = {ws.cell(row=r, column=2).value: r for r in range(13, 93) if ws.cell(row=r, column=2).value}
OUT = {"final": "C6", "dec": "C7", "cls": "C8", "malus": "C12", "inc": "E23", "eco": "C10", "guar": "C16", "econ": "C15",
       "d1": "C32", "d2": "C33", "d3": "C34", "d4": "C35", "miss": "C24", "alr": "C26", "adj": "C11", "coef": "C14", "unk": "E19"}
def run(case):
    inp = {}
    for k, r in rowof.items():
        v = case["inputs"].get(k)
        if isinstance(v, bool): v = "Oui" if v else "Non"
        inp[ref("Saisie", f"E{r}")] = v if v is not None else formulas.functions.EMPTY if hasattr(formulas.functions, "EMPTY") else ""
    inp[ref("Saisie", "C6")] = case["segment"] or ""
    inp[ref("Saisie", "C7")] = case["zone"] or ""
    inp[ref("Saisie", "C8")] = case["cls"] or ""
    sol = xl.calculate(inputs=inp, outputs=[ref("Resultat", c) for c in OUT.values()])
    res = {}
    for name, c in OUT.items():
        v = sol[ref("Resultat", c)].value
        v = v[0][0] if hasattr(v, "__getitem__") else v
        res[name] = v
    return res
sel = sys.argv[1:] 
ok = 0; tot = 0; report = []
for case in vec:
    if sel and case["id"] not in sel: continue
    t = time.time(); r = run(case); e = case["expected"]; tot += 1
    def num(x):
        try: return float(x)
        except: return float('nan')
    checks = {
      "final": abs(num(r["final"]) - e["scoreFinal"]) <= 0.02, "dec": str(r["dec"]) == e["decision"], "cls": str(r["cls"]) == e["internalClass"],
      "malus": abs(num(r["malus"]) - e["totalMalus"]) <= 0.001, "inc": (str(r["inc"]).upper() in ("TRUE", "1")) == e["dataIncomplete"],
      "eco": abs(num(r["eco"]) - e["scoreEco"]) <= 0.02, "d1": abs(num(r["d1"]) - e["domains"]["D1"]) <= 0.02, "d2": abs(num(r["d2"]) - e["domains"]["D2"]) <= 0.02,
      "d3": abs(num(r["d3"]) - e["domains"]["D3"]) <= 0.02, "d4": abs(num(r["d4"]) - e["domains"]["D4"]) <= 0.02,
      "econ": abs(num(r["econ"]) - e["economicScore"]) <= 0.02,
      "guar": (str(r["guar"]) == "n/a") if e["guaranteeScore"] is None else abs(num(r["guar"]) - e["guaranteeScore"]) <= 0.02,
    }
    good = all(checks.values()); ok += good
    if sel: print("    brut:", {k: (type(v).__name__, str(v)[:60]) for k, v in r.items()})
    print(f"{case['id']} {'OK ' if good else 'ÉCART'} {time.time()-t:5.1f}s  final={r['final']} (att {e['scoreFinal']}) {r['dec']} (att {e['decision']})", flush=True)
    if not good: print("    écarts:", [k for k, v in checks.items() if not v], "| obtenu:", {k: r[k] for k in r}, flush=True)
print(f"\nBILAN: {ok}/{tot} cas conformes")
